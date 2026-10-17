"""
Generate Architecture Diagram as PowerPoint (.pptx)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_arch_pptx(output_path="presentation/ARCHITECTURE_DIAGRAM.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Colors
    BG = RGBColor(11, 17, 32)
    CARD = RGBColor(30, 41, 59)
    C_BLUE = RGBColor(56, 189, 248)
    C_PURPLE = RGBColor(168, 85, 247)
    C_GREEN = RGBColor(52, 211, 153)
    C_AMBER = RGBColor(251, 191, 36)
    TXT_LIGHT = RGBColor(241, 245, 249)
    TXT_MUTED = RGBColor(148, 163, 184)
    
    # BG
    bg_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_box.fill.solid()
    bg_box.fill.fore_color.rgb = BG
    bg_box.line.fill.background()
    
    # Title
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.133), Inches(0.8))
    tf = tb.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "GlycoTwin-AI: End-to-End Multimodal Digital Twin Architecture"
    p1.font.bold = True
    p1.font.size = Pt(20)
    p1.font.color.rgb = TXT_LIGHT
    
    p2 = tf.add_paragraph()
    p2.text = "Happiest Health Reimagining & Reforming Healthcare Summit 2026 • Proactive Diabetes Prevention in India"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TXT_MUTED
    
    pillars = [
        ("1. DUAL DATA STREAMS", "DPDP Act 2023 & HIPAA Sandbox", C_BLUE, Inches(0.6), [
            "Static Historical EHR:",
            "• 1,000 Synthea Indian Profiles",
            "• Asian Indian BMI (>=23 kg/m²)",
            "• Baseline HbA1c (5.7 - 14.2%)",
            "• Fasting Blood Glucose (FBG)",
            "• eGFR & Serum Creatinine",
            "• Triglycerides / HDL Profile",
            "• Comorbidities & TCF7L2 gene",
            "• Baseline Pharmacotherapy",
            "",
            "Dynamic IoT Telemetry:",
            "• 60,120 Timepoints (5-min)",
            "• Interstitial Glucose (CGM)",
            "• Continuous Heart Rate (PPG)",
            "• HRV (RMSSD Autonomic tone)",
            "• Step Cadence & Sleep Stages",
            "• Indian Meal Carbs & Insulins"
        ]),
        ("2. MULTIMODAL FUSION", "Cross-Domain Feature Synthesis", C_PURPLE, Inches(3.7), [
            "Temporal Dynamic Kinetics:",
            "• Glucose Velocity dG/dt (15/30m)",
            "• Glycemic Acceleration d²G/dt²",
            "• Rolling Mean & Std (1h / 3h)",
            "• Intraday CV% Volatility",
            "",
            "Autonomic Coupling:",
            "• Autonomic Stress: HR / HRV",
            "• Pre-hypo Sympathetic Surge",
            "• 1-Hour Step Accumulation",
            "• 2-Hour Carbohydrate Load",
            "",
            "Latent Patient Embedding:",
            "• Asian Indian Metabolic Profile",
            "• Trigonometric Circadian Wave",
            "• Beta-Cell Reserve Proxy",
            "★ 42 Multimodal Features"
        ]),
        ("3. DUAL-HEAD AI CORE", "Prediction & In-Silico Engine", C_GREEN, Inches(6.8), [
            "Head 1: 2h Glucose Forecaster",
            "• HistGBM Multi-Tree Regressor",
            "• 2-Hour Horizon (+120 mins)",
            "★ MARD: 7.15% (Gold Std <10%)",
            "★ RMSE: 14.50 mg/dL",
            "• Dynamic 95% Confidence Band",
            "",
            "Head 2: Adverse Classifier",
            "• Severe Postprandial Hyper (>200)",
            "  Precision = 65.31%",
            "• Nocturnal Hypoglycemia (<70)",
            "  Early Warning: 120 mins",
            "",
            "In-Silico What-If Simulator",
            "• Post-Meal Walk (0-45 mins)",
            "• Carb Swap (-10 to -60g)",
            "• Basal Titration (-4U to +4U)"
        ]),
        ("4. CLINICIAN COCKPIT", "Interactive Doctor Dashboard", C_AMBER, Inches(9.9), [
            "Clinical Web Interface:",
            "• Virtual Patient Selector",
            "• Static EHR Profile Cards",
            "• 6h Multi-stream Waveforms",
            "• 2h Forward Projection Band",
            "• Dynamic Triage Risk Light",
            "",
            "Physician Interactivity:",
            "• Real-time What-If Sliders",
            "• Dynamic Delta Metric Output",
            "• In-Silico Trajectory Overlay",
            "• RSSDI & ADA Decision Support",
            "",
            "Explainable AI (XAI):",
            "• Physiological Root-Cause XAI",
            "• Autonomic Crash Alerts",
            "★ FastAPI + Chart.js Web App"
        ])
    ]
    
    for title, sub, color, left, items in pillars:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.3), Inches(2.8), Inches(5.6))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.color.rgb = color
        card.line.width = Pt(1.5)
        
        ptb = slide.shapes.add_textbox(left + Inches(0.15), Inches(1.4), Inches(2.5), Inches(5.4))
        ptf = ptb.text_frame
        ptf.word_wrap = True
        ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0
        
        p = ptf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = color
        
        ps = ptf.add_paragraph()
        ps.text = sub
        ps.font.size = Pt(8.5)
        ps.font.color.rgb = TXT_MUTED
        ps.space_after = Pt(8)
        
        for it in items:
            pi = ptf.add_paragraph()
            pi.text = it
            if it.startswith("★"):
                pi.font.bold = True
                pi.font.size = Pt(9.5)
                pi.font.color.rgb = C_GREEN
            elif ":" in it and not it.startswith("•"):
                pi.font.bold = True
                pi.font.size = Pt(9.5)
                pi.font.color.rgb = TXT_LIGHT
            else:
                pi.font.size = Pt(8.5)
                pi.font.color.rgb = TXT_MUTED
                
    prs.save(output_path)
    print(f"✓ Saved Architecture PPTX: {output_path}")

if __name__ == "__main__":
    create_arch_pptx()
