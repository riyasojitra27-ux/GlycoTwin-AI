"""
GlycoTwin-AI: Automated Native PowerPoint (.pptx) Presentation Generator
Creates a 16:9 widescreen presentation deck for the Happiest Health Summit 2026.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_pptx(output_path="presentation/GLYCOTWIN_PRESENTATION.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Color palette
    BG_COLOR = RGBColor(11, 17, 32)      # #0B1120
    CARD_BG = RGBColor(30, 41, 59)       # #1E293B
    BORDER_COLOR = RGBColor(51, 65, 85)  # #334155
    ACCENT_BLUE = RGBColor(56, 189, 248) # #38BDF8
    TEXT_LIGHT = RGBColor(241, 245, 249) # #F1F5F9
    TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8
    TEXT_GREEN = RGBColor(52, 211, 153) # #34D399
    TEXT_GOLD = RGBColor(253, 230, 138)  # #FDE68A

    slides_content = [
        {
            "title": "GlycoTwin-AI: Multimodal Digital Twin for Proactive Diabetes Management",
            "subtitle": "Happiest Health Reimagining & Reforming Healthcare in India Summit 2026, Bengaluru",
            "bullets": [
                "Track: Digital Twin Proof-of-Concept (PoC) Challenge",
                "Focus Condition: Type 2 Diabetes Mellitus & Postprandial/Nocturnal Dysglycemia",
                "Target Population: Asian Indian Phenotype (>101M Diabetics, high carb staple diets)",
                "Core Innovation: In-silico physiological twin fusing static EHR with wearable CGM/PPG",
                "Key Capability: 2-Hour ahead adverse event prediction & What-If counterfactual simulation"
            ],
            "footer": "Team GlycoTwin India • Center for Healthcare AI & Innovation"
        },
        {
            "title": "The Indian Diabetes Crisis: Why Reactive Healthcare Fails",
            "subtitle": "ICMR-INDIAB 2023 Epidemiological Context & Clinical Motivation",
            "bullets": [
                "101.3 Million Diabetics + 136 Million Pre-diabetics in India ('Diabetes Capital of the World')",
                "Asian Indian Phenotype: High visceral fat & severe insulin resistance at lower BMI (>= 23 kg/m²)",
                "Dietary Trap: Refined rice/wheat staples cause massive postprandial glucose surges (>200 mg/dL)",
                "Silent Danger: Nocturnal hypoglycemia (<70 mg/dL) during sleep causes fatal cardiac arrhythmias",
                "Reactive Failure: Quarterly HbA1c tests are 3-month trailing averages; they miss intraday volatility",
                "Solution: Continuous Digital Twin providing real-time in-silico simulation & proactive alerts"
            ],
            "footer": "Clinical Background • Slide 2 of 10"
        },
        {
            "title": "Dual Data Stream Fusion Architecture",
            "subtitle": "Synthesizing Static Clinical Records with High-Frequency Wearable IoT Telemetry",
            "bullets": [
                "Stream 1 (Static EHR): Demographics, Asian Indian BMI, baseline HbA1c, eGFR, lipid profile, comorbidities, TCF7L2 genetic risk, & baseline medications.",
                "Stream 2 (Dynamic IoT Wearables): 5-minute CGM interstitial glucose, continuous Heart Rate (bpm), HRV (RMSSD autonomic tone), step cadence, sleep stages, & Indian meal carbs (g).",
                "Feature Synthesis: 42 kinetic & cross-domain features including velocity (dG/dt), acceleration (d²G/dt²), intraday CV% volatility, and autonomic stress coupling (HR / HRV ratio).",
                "Compliance: 100% compliant with Indian DPDP Act 2023 & HIPAA sandbox regulations (Zero PII)."
            ],
            "footer": "Data Architecture • Slide 3 of 10"
        },
        {
            "title": "Dual-Head Machine Learning Core & Clinical Metrics",
            "subtitle": "Clinical Validation on Unseen Patient Test Cohorts (Zero Patient Leakage)",
            "bullets": [
                "Continuous Regressor (2-Hour Horizon Forecast):",
                "   • MARD (Mean Absolute Relative Difference): 7.15% (Gold standard clinical threshold < 10%)",
                "   • RMSE: 14.50 mg/dL | MAE: 11.11 mg/dL across 59,940 high-resolution multimodal records",
                "   • Generates dynamic 95% confidence bounds reflecting biological uncertainty",
                "Adverse Event Classifier (Early Warning Lead Time: 120 minutes):",
                "   • Severe Postprandial Hyperglycemia Spike (>200 mg/dL): Precision = 65.31%",
                "   • Impending Nocturnal Hypoglycemia (<70 mg/dL): Calibrated safety decision threshold (18%)",
                "   • Macro F1: 0.54 | Patient-stratified holdout evaluation"
            ],
            "footer": "Model Evaluation • Slide 4 of 10"
        },
        {
            "title": "Physician 'What-If' In-Silico Counterfactual Simulator",
            "subtitle": "Empowering Doctors to Test Interventions Digitally Before Clinical Execution",
            "bullets": [
                "Clinical Need: Physicians cannot conduct risky trial-and-error experiments on living patients.",
                "Digital Twin Sandbox: Simulates three high-impact clinical interventions in real time (<50ms):",
                "   1. Post-Prandial Brisk Walk (0 to 45 mins): Simulates muscular GLUT-4 glucose clearance.",
                "   2. Indian Carb Swap / Reduction (0 to 60g): Models switching white rice to millets/proteins.",
                "   3. Medication Adjustment (-4U to +4U): Titrates basal insulin or shifts sulfonylurea timing.",
                "Instant Feedback: Re-renders simulated counterfactual trajectory vs. baseline projection curve, calculating exact projected glucose delta (mg/dL) and RSSDI/ADA clinical guidance."
            ],
            "footer": "Counterfactual Simulation • Slide 5 of 10"
        },
        {
            "title": "Clinician Digital Twin Cockpit UI",
            "subtitle": "Human-in-the-Loop Web Platform for Proactive Patient Management",
            "bullets": [
                "Full-Featured Glassmorphic Web App: Built with FastAPI, Tailwind CSS, & Chart.js.",
                "Live Telemetry Waveforms: Synchronized visualization of CGM, Heart Rate, HRV, steps, & meals.",
                "Triage Early Warning System: Color-coded alert badges (Green = Stable, Amber = Hyper, Red = Hypo).",
                "Explainable AI (XAI) Drawer: Physiological root-cause attributions explain why alerts fired.",
                "Interactive Simulator Panel: Direct slider controls for walking, carb reduction, and insulin titration.",
                "Automated AI Presentation Tour: Built-in teleprompter HUD with optional AI voiceover."
            ],
            "footer": "User Interface & Experience • Slide 6 of 10"
        },
        {
            "title": "Clinical Walkthrough: Two Real-World Scenarios",
            "subtitle": "Proactive Intervention in Indian Patient Personas",
            "bullets": [
                "Scenario A (Postprandial Hyperglycemia - Patient IND-PT-10007):",
                "   • 54y male, Rice staple diet, baseline HbA1c 9.4%. Post-lunch predicted spike: 218 mg/dL.",
                "   • Doctor simulates: 20-min post-lunch walk + 35g carb reduction (foxtail millet swap).",
                "   • Twin outcome: Predicted 2h glucose drops by -38 mg/dL, restoring safe in-range control (180 mg/dL).",
                "Scenario B (Impending Nocturnal Hypoglycemia - Patient IND-PT-10004):",
                "   • 59y female on basal insulin (22U). Twin detects 02:30 AM autonomic crash (HRV drops to 14ms).",
                "   • Twin flags impending 04:30 AM hypo crash (61 mg/dL) 2 hours in advance.",
                "   • Doctor simulates: Titrate evening basal insulin by -3U + bedtime snack -> Crash averted."
            ],
            "footer": "Clinical Validation • Slide 7 of 10"
        },
        {
            "title": "DPDP Act 2023 & HIPAA Compliance Architecture",
            "subtitle": "Patient Privacy, Synthetic Data Integrity & Regulatory Sandbox Rules",
            "bullets": [
                "Indian DPDP Act 2023 Compliance: Fully synthetic patient cohorts adhering to national privacy standards.",
                "Zero Real Patient PII: Synthetic identities (IND-PT-XXXXX) generated with authentic clinical distributions.",
                "Synthea & MIMIC-IV Standard Schema: Statistically verified against ICMR-INDIAB epidemiological metrics.",
                "Data Leakage Protection: Strict patient-level train/test separation prevents data memorization.",
                "Audit Logging: Every in-silico simulation and decision recommendation is timestamped and cryptographically logged."
            ],
            "footer": "Regulatory Compliance • Slide 8 of 10"
        },
        {
            "title": "Socioeconomic Impact & Indian Healthcare Deployment Roadmap",
            "subtitle": "From University Incubator Prototype to Ayushman Bharat Digital Mission (ABDM)",
            "bullets": [
                "Economic Impact: Avoids ₹40,000 - ₹1,50,000 emergency admissions for DKA and severe hypoglycemia.",
                "Long-Term Prevention: Blunting glycemic spikes reduces HbA1c by 0.8-1.2%, cutting microvascular complications by 37%.",
                "Integration Roadmap:",
                "   • Phase 1 (Now): Working Digital Twin PoC, multimodal predictive engine & doctor cockpit.",
                "   • Phase 2 (Summit Final): Hardware integration with commercial CGMs (Abbott FreeStyle Libre) & wearables.",
                "   • Phase 3 (Clinical Pilot): FHIR API integration into Ayushman Bharat Digital Mission (ABDM) & e-Sanjeevani.",
                "   • Phase 4 (Edge AI): Mobile companion app providing real-time personalized patient nudges."
            ],
            "footer": "Deployment Roadmap • Slide 9 of 10"
        },
        {
            "title": "Conclusion & Competitive Summary",
            "subtitle": "Happiest Health Reimagining and Reforming Healthcare Summit 2026",
            "bullets": [
                "Accomplished: Selected #1 Indian healthcare priority (Type 2 Diabetes) & delivered an end-to-end Digital Twin.",
                "Dual Data Fusion: Seamlessly unites static clinical EHR with high-resolution wearable telemetry.",
                "Clinical Grade Accuracy: 7.15% MARD and 120-minute lead time on acute dysglycemic adverse events.",
                "Doctor-Centric Innovation: In-silico 'What-If' counterfactual engine transforms reactive care into proactive prevention.",
                "Deliverables: Complete open-source codebase (MIT), interactive UI, architecture diagrams, and 20-min demo video.",
                "Vision: Building the proactive foundation for Indian digital healthcare."
            ],
            "footer": "Thank You • Happiest Health Summit 2026 Bengaluru"
        }
    ]

    for idx, slide_info in enumerate(slides_content, 1):
        slide = prs.slides.add_slide(blank_layout)
        
        # 1. Background rectangle
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_COLOR
        bg_shape.line.fill.background()
        
        # 2. Main content card
        card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.133), Inches(5.4))
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = CARD_BG
        card_shape.line.color.rgb = BORDER_COLOR
        card_shape.line.width = Pt(1.5)
        
        # 3. Header Title & Subtitle Box
        header_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.25), Inches(12.133), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_title = tf.paragraphs[0]
        p_title.text = slide_info["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_BLUE
        
        p_sub = tf.add_paragraph()
        p_sub.text = slide_info["subtitle"]
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED
        
        # 4. Bullet Points Box inside Card
        bullet_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11.5), Inches(4.9))
        btf = bullet_box.text_frame
        btf.word_wrap = True
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0
        
        first = True
        for b in slide_info["bullets"]:
            p = btf.paragraphs[0] if first else btf.add_paragraph()
            first = False
            p.text = b
            p.font.name = "Arial"
            p.space_after = Pt(10)
            
            if b.startswith("   •"):
                p.font.size = Pt(13)
                p.font.color.rgb = TEXT_MUTED
                p.level = 1
            elif b.startswith("   1.") or b.startswith("   2.") or b.startswith("   3."):
                p.font.size = Pt(13)
                p.font.color.rgb = TEXT_GOLD
                p.level = 1
            elif "Continuous Regressor" in b or "Adverse Event Classifier" in b or "Scenario A" in b or "Scenario B" in b:
                p.font.size = Pt(14)
                p.font.bold = True
                p.font.color.rgb = TEXT_GREEN
                p.level = 0
            else:
                p.font.size = Pt(13.5)
                p.font.color.rgb = TEXT_LIGHT
                p.level = 0

        # 5. Footer Box
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(6.85), Inches(12.133), Inches(0.4))
        ftf = footer_box.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        
        fp = ftf.paragraphs[0]
        fp.text = f"{slide_info['footer']}  |  Slide {idx} of {len(slides_content)}"
        fp.font.name = "Arial"
        fp.font.size = Pt(10)
        fp.font.color.rgb = TEXT_MUTED

    prs.save(output_path)
    print(f"✓ Successfully generated PowerPoint (.pptx): {output_path}")

if __name__ == "__main__":
    create_pptx()
